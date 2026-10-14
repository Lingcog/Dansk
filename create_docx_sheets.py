import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_danish_doc(output_path):
    doc = docx.Document()
    
    # Page setup - Margins (0.5 in / 1.27 cm for crisp A4 fit)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("UDTALE-WORKSHOP")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = RGBColor(192, 0, 0) # Red
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    run_sub = p_sub.add_run("Makker-Lytteark (Peer Listening) — Træning af konsonanterne D, R og G (DU3 Modul 1-2)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    p_sub2 = doc.add_paragraph()
    p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub2.paragraph_format.space_before = Pt(0)
    p_sub2.paragraph_format.space_after = Pt(8)
    run_sub2 = p_sub2.add_run("Samtale 1: En ansat ringer til lægen")
    run_sub2.font.name = "Arial"
    run_sub2.font.size = Pt(11)
    run_sub2.bold = True
    run_sub2.font.color.rgb = RGBColor(50, 50, 50)
    
    # Red divider bar
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(8)
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="C00000"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(p_line_border)

    # Student Names Table
    name_table = doc.add_table(rows=1, cols=2)
    name_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    name_table.autofit = False
    name_table.columns[0].width = Inches(3.5)
    name_table.columns[1].width = Inches(3.5)
    
    c0 = name_table.cell(0, 0)
    p_c0 = c0.paragraphs[0]
    p_c0.paragraph_format.space_after = Pt(6)
    r_n1 = p_c0.add_run("Kursistens navn (taler): ")
    r_n1.bold = True
    r_n1.font.size = Pt(9.5)
    p_c0.add_run("___________________________")
    
    c1 = name_table.cell(0, 1)
    p_c1 = c1.paragraphs[0]
    p_c1.paragraph_format.space_after = Pt(6)
    r_n2 = p_c1.add_run("Makkers navn (lytter): ")
    r_n2.bold = True
    r_n2.font.size = Pt(9.5)
    p_c1.add_run("___________________________")
    
    # Instructions Box
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(4)
    p_inst.paragraph_format.space_after = Pt(8)
    # Left border and light background
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="0078D4"/></w:pBdr>')
    p_inst._p.get_or_add_pPr().append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F8FC"/>')
    p_inst._p.get_or_add_pPr().append(shd)
    
    r_i_title = p_inst.add_run("INSTRUKTION TIL LYTTEREN:\n")
    r_i_title.bold = True
    r_i_title.font.size = Pt(9.5)
    r_i_title.font.color.rgb = RGBColor(0, 120, 212)
    
    r_i_body = p_inst.add_run(
        "1. Din makker læser replikkerne højt (eller bruger appen).\n"
        "2. Lyt opmærksomt efter de markerede lyde for D, R og G.\n"
        "3. Sæt kryds: Hørte du den rigtige danske lyd, eller skal den justeres?"
    )
    r_i_body.font.size = Pt(9)
    r_i_body.font.color.rgb = RGBColor(40, 40, 40)
    
    # Main Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    table.columns[0].width = Inches(3.9)
    table.columns[1].width = Inches(1.6)
    table.columns[2].width = Inches(1.6)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ["REPLIKSKIFTE & FOKUSORD", "RIGTIGT", "FORKERT"]
    hdr_colors = [RGBColor(40, 40, 40), RGBColor(16, 124, 65), RGBColor(192, 0, 0)]
    
    for idx, (cell, title, color) in enumerate(zip(hdr_cells, hdr_titles, hdr_colors)):
        set_cell_background(cell, "F2F4F7")
        set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = color
    
    items = [
        (
            '1. "Hej, jeg hedder [Navn]."',
            [
                ('• Stumt G i "jeg" (udtales som [ja])', '☐ G var stumt!', '☐ Hårdt G blev udtalt'),
                ('• Blødt D i "hedder" (tungen bag tænderne)', '☐ Fint blødt D!', '☐ Hårdt D som engelsk'),
                ('• Vokalisk R i "hedder" (åben -er endelse)', '☐ Blød, åben lyd', '☐ Hårdt / rullende R')
            ]
        ),
        (
            '2. "Jeg ringer fordi jeg har det virkelig dårligt."',
            [
                ('• Stumt G i "jeg" [ja] og "virkelig"', '☐ G var stumt!', '☐ Hårdt G blev udtalt'),
                ('• Hårdt hals-R i "ringer" (ganelyd)', '☐ Godt hals-R!', '☐ Rullende tungespids-R'),
                ('• Vokalisk R i "har" (lang, åben vokal)', '☐ Vokalen forlænges', '☐ Hårdt R udtalt')
            ]
        ),
        (
            '3. "Jeg har høj feber, hovedpine og muskelsmerter."',
            [
                ('• Vokalisk R i "feber", "har" og "muskelsmerter"', '☐ Blød -er endelse', '☐ Hård konsonant'),
                ('• Blødt D i "hovedpine" (blød luftlyd)', '☐ Flot blødt D!', '☐ Hårdt D (d-lyd)')
            ]
        ),
        (
            '4. "Ja, det gør jeg."',
            [
                ('• Vokalisk R i "gør" (smelter med vokalen)', '☐ Vokalisk R', '☐ Hårdt R'),
                ('• Stumt G i "jeg" (udtales [ja])', '☐ G var stumt!', '☐ G blev udtalt')
            ]
        ),
        (
            '5. "Ja, der er lidt grønt slim."',
            [
                ('• Hårdt G og hårdt hals-R i klyngen "grønt"', '☐ Tydeligt gr-!', '☐ Svag konsonantklynge'),
                ('• Vokalisk R i "der" og "er" (åbne vokaler)', '☐ Smelter med vokal', '☐ For hårdt R')
            ]
        ),
        (
            '6. "Nej, kun træthed og ondt i halsen."',
            [
                ('• Hårdt hals-R i "træthed" (efter t)', '☐ Flot tr- klynge', '☐ Tungespids-R'),
                ('• Blødt D i "træthed" (-hed endelse)', '☐ Blødt D [ð]', '☐ Hårdt D'),
                ('• Stumt D i "ondt" (-nd klynge efter n)', '☐ D var stumt!', '☐ D blev udtalt')
            ]
        ),
        (
            '7. "Tak, det vil jeg gøre."',
            [
                ('• Stumt G i "jeg" [ja]', '☐ G var stumt!', '☐ G blev udtalt'),
                ('• Vokalisk R i "gøre" (åben -re endelse)', '☐ Blød, åben lyd', '☐ Hårdt konsonant-R')
            ]
        )
    ]

    for turn_idx, (turn_title, bullets) in enumerate(items):
        row = table.add_row()
        cells = row.cells
        
        # Alternating background
        bg = "FFFFFF" if turn_idx % 2 == 0 else "FAFAFA"
        for c in cells:
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        
        # Column 0: Title + Bullet points
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r_t = p0.add_run(turn_title + "\n")
        r_t.bold = True
        r_t.font.size = Pt(9)
        r_t.font.color.rgb = RGBColor(20, 20, 20)
        
        # Add bullets in col 0, col 1, col 2
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.add_run("\n") # to align with title
        
        p2 = cells[2].paragraphs[0]
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.add_run("\n")
        
        for b_text, c_correct, c_incorrect in bullets:
            r_b = p0.add_run(f"{b_text}\n")
            r_b.font.size = Pt(8.5)
            r_b.font.color.rgb = RGBColor(60, 60, 60)
            
            r_c1 = p1.add_run(f"{c_correct}\n")
            r_c1.font.size = Pt(8.5)
            r_c1.font.color.rgb = RGBColor(16, 124, 65)
            
            r_c2 = p2.add_run(f"{c_incorrect}\n")
            r_c2.font.size = Pt(8.5)
            r_c2.font.color.rgb = RGBColor(180, 0, 0)
            
    # Bottom Summary Box
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(8)
    p_sum.paragraph_format.space_after = Pt(2)
    bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="12" w:space="4" w:color="C00000"/></w:pBdr>')
    p_sum._p.get_or_add_pPr().append(bdr)
    
    r_s_title = p_sum.add_run("OPSAMLING & FEEDBACK TIL DIN MAKKER:\n")
    r_s_title.bold = True
    r_s_title.font.size = Pt(9.5)
    r_s_title.font.color.rgb = RGBColor(192, 0, 0)
    
    r_s_body = p_sum.add_run(
        "Hvilket bogstav skal din makker have lidt ekstra fokus på i appen?\n"
        "☐ Bogstav D (blødt D / stumt D)         "
        "☐ Bogstav R (hals-R / vokalisk R)         "
        "☐ Bogstav G (stumt G / hårdt G)"
    )
    r_s_body.font.size = Pt(9)
    r_s_body.font.color.rgb = RGBColor(40, 40, 40)
    
    doc.save(output_path)
    print("Saved Danish docx to:", output_path)

def create_english_doc(output_path):
    doc = docx.Document()
    
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("PRONUNCIATION WORKSHOP")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = RGBColor(192, 0, 0)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    run_sub = p_sub.add_run("Peer Listening Sheet — Training the consonants D, R and G (DU3 Module 1-2)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    p_sub2 = doc.add_paragraph()
    p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub2.paragraph_format.space_before = Pt(0)
    p_sub2.paragraph_format.space_after = Pt(8)
    run_sub2 = p_sub2.add_run("Dialogue 1: An Employee Calls the Doctor")
    run_sub2.font.name = "Arial"
    run_sub2.font.size = Pt(11)
    run_sub2.bold = True
    run_sub2.font.color.rgb = RGBColor(50, 50, 50)
    
    # Red divider bar
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(8)
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="C00000"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(p_line_border)

    # Student Names Table
    name_table = doc.add_table(rows=1, cols=2)
    name_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    name_table.autofit = False
    name_table.columns[0].width = Inches(3.5)
    name_table.columns[1].width = Inches(3.5)
    
    c0 = name_table.cell(0, 0)
    p_c0 = c0.paragraphs[0]
    p_c0.paragraph_format.space_after = Pt(6)
    r_n1 = p_c0.add_run("Speaker's Name (Student): ")
    r_n1.bold = True
    r_n1.font.size = Pt(9.5)
    p_c0.add_run("___________________________")
    
    c1 = name_table.cell(0, 1)
    p_c1 = c1.paragraphs[0]
    p_c1.paragraph_format.space_after = Pt(6)
    r_n2 = p_c1.add_run("Listener's Name (Partner): ")
    r_n2.bold = True
    r_n2.font.size = Pt(9.5)
    p_c1.add_run("___________________________")
    
    # Instructions Box
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(4)
    p_inst.paragraph_format.space_after = Pt(8)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="0078D4"/></w:pBdr>')
    p_inst._p.get_or_add_pPr().append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F8FC"/>')
    p_inst._p.get_or_add_pPr().append(shd)
    
    r_i_title = p_inst.add_run("INSTRUCTIONS FOR THE LISTENER:\n")
    r_i_title.bold = True
    r_i_title.font.size = Pt(9.5)
    r_i_title.font.color.rgb = RGBColor(0, 120, 212)
    
    r_i_body = p_inst.add_run(
        "1. Your partner reads the Danish sentences aloud (or uses the app).\n"
        "2. Listen carefully for the targeted Danish consonant sounds (D, R, G).\n"
        "3. Check the box for what you actually hear: Correct Danish sound vs. needs practice!"
    )
    r_i_body.font.size = Pt(9)
    r_i_body.font.color.rgb = RGBColor(40, 40, 40)
    
    # Main Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    table.columns[0].width = Inches(3.9)
    table.columns[1].width = Inches(1.6)
    table.columns[2].width = Inches(1.6)
    
    hdr_cells = table.rows[0].cells
    hdr_titles = ["TURNS & TARGET WORDS", "CORRECT", "NEEDS PRACTICE"]
    hdr_colors = [RGBColor(40, 40, 40), RGBColor(16, 124, 65), RGBColor(192, 0, 0)]
    
    for idx, (cell, title, color) in enumerate(zip(hdr_cells, hdr_titles, hdr_colors)):
        set_cell_background(cell, "F2F4F7")
        set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = color
    
    items_en = [
        (
            '1. "Hej, jeg hedder [Name]."',
            [
                ('• Silent G in "jeg" (pronounced like [ya])', '☐ G was silent!', '☐ Hard G pronounced'),
                ('• Soft D in "hedder" (tongue flat behind teeth)', '☐ Great soft D!', '☐ Hard English D'),
                ('• Vowel R in "hedder" (open relaxed -er ending)', '☐ Soft vowel ending', '☐ Hard / rolled R')
            ]
        ),
        (
            '2. "Jeg ringer fordi jeg har det virkelig dårligt."',
            [
                ('• Silent G in "jeg" [ya] and "virkelig"', '☐ G was silent!', '☐ Hard G pronounced'),
                ('• Throat R in "ringer" (gargling sound in throat)', '☐ Good throat R!', '☐ Rolled tongue-tip R'),
                ('• Vowel R in "har" (lengthened open vowel)', '☐ Vowel lengthened', '☐ Hard R pronounced')
            ]
        ),
        (
            '3. "Jeg har høj feber, hovedpine og muskelsmerter."',
            [
                ('• Vowel R in "feber", "har" and "muskelsmerter"', '☐ Relaxed -er sound', '☐ Sharp consonant R'),
                ('• Soft D in "hovedpine" (like "th" in "weather")', '☐ Great soft D!', '☐ Hard English D')
            ]
        ),
        (
            '4. "Ja, det gør jeg."',
            [
                ('• Vowel R in "gør" (melts into the vowel)', '☐ Smooth vowel', '☐ Hard throat R'),
                ('• Silent G in "jeg" (pronounced [ya])', '☐ G was silent!', '☐ G was pronounced')
            ]
        ),
        (
            '5. "Ja, der er lidt grønt slim."',
            [
                ('• Hard G + throat R cluster in "grønt"', '☐ Clear gr- sound!', '☐ Weak cluster'),
                ('• Vowel R in "der" and "er" (open vowel quality)', '☐ Melts with vowel', '☐ Too harsh R sound')
            ]
        ),
        (
            '6. "Nej, kun træthed og ondt i halsen."',
            [
                ('• Throat R in "træthed" (tight cluster after t)', '☐ Good tr- sound!', '☐ Front / rolled R'),
                ('• Soft D in "træthed" (-hed ending)', '☐ Soft D [ð]', '☐ Hard D sound'),
                ('• Silent D in "ondt" (-nd cluster after n)', '☐ D was silent!', '☐ D was pronounced')
            ]
        ),
        (
            '7. "Tak, det vil jeg gøre."',
            [
                ('• Silent G in "jeg" [ya]', '☐ G was silent!', '☐ G was pronounced'),
                ('• Vowel R in "gøre" (relaxed open ending)', '☐ Soft open ending', '☐ Sharp consonant R')
            ]
        )
    ]

    for turn_idx, (turn_title, bullets) in enumerate(items_en):
        row = table.add_row()
        cells = row.cells
        
        bg = "FFFFFF" if turn_idx % 2 == 0 else "FAFAFA"
        for c in cells:
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r_t = p0.add_run(turn_title + "\n")
        r_t.bold = True
        r_t.font.size = Pt(9)
        r_t.font.color.rgb = RGBColor(20, 20, 20)
        
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.add_run("\n")
        
        p2 = cells[2].paragraphs[0]
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.add_run("\n")
        
        for b_text, c_correct, c_incorrect in bullets:
            r_b = p0.add_run(f"{b_text}\n")
            r_b.font.size = Pt(8.5)
            r_b.font.color.rgb = RGBColor(60, 60, 60)
            
            r_c1 = p1.add_run(f"{c_correct}\n")
            r_c1.font.size = Pt(8.5)
            r_c1.font.color.rgb = RGBColor(16, 124, 65)
            
            r_c2 = p2.add_run(f"{c_incorrect}\n")
            r_c2.font.size = Pt(8.5)
            r_c2.font.color.rgb = RGBColor(180, 0, 0)
            
    # Bottom Summary Box
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(8)
    p_sum.paragraph_format.space_after = Pt(2)
    bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="12" w:space="4" w:color="C00000"/></w:pBdr>')
    p_sum._p.get_or_add_pPr().append(bdr)
    
    r_s_title = p_sum.add_run("SUMMARY & FEEDBACK FOR YOUR PARTNER:\n")
    r_s_title.bold = True
    r_s_title.font.size = Pt(9.5)
    r_s_title.font.color.rgb = RGBColor(192, 0, 0)
    
    r_s_body = p_sum.add_run(
        "Which consonant should your partner practice more in the app?\n"
        "☐ Letter D (Soft D / Silent D)         "
        "☐ Letter R (Throat R / Vowel R)         "
        "☐ Letter G (Silent G / Hard G)"
    )
    r_s_body.font.size = Pt(9)
    r_s_body.font.color.rgb = RGBColor(40, 40, 40)
    
    doc.save(output_path)
    print("Saved English docx to:", output_path)

if __name__ == '__main__':
    create_danish_doc('public/docs/Peer_Listening_Sheet_Samtale1_Dansk.docx')
    create_english_doc('public/docs/Peer_Listening_Sheet_Samtale1_English.docx')
