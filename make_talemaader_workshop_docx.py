import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins: 2 cm (~0.79 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    # Colors
    NAVY = RGBColor(30, 58, 138)       # #1e3a8a
    TEAL = RGBColor(13, 148, 136)      # #0d9488
    DARK_GRAY = RGBColor(71, 85, 105)  # #475569
    BLACK = RGBColor(15, 23, 42)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("DANSKUDDANNELSE 3 (DU3) • WORKSHOP & UNDERVISNINGSMATERIALE\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = TEAL

    run_title = title_p.add_run("Livets gang med danske talemåder")
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(14)
    run_desc = subtitle_p.add_run(
        "Stilladseret opgavesæt: Kollaborativ afkodning i appen, betydningsvalg, livssituationer og A/B rollespil med udtale og samtaleteknik."
    )
    run_desc.font.size = Pt(11)
    run_desc.font.italic = True
    run_desc.font.color.rgb = DARK_GRAY

    # Overview box
    box_table = doc.add_table(rows=1, cols=1)
    box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = box_table.cell(0, 0)
    set_cell_background(c, "F1F5F9")
    set_cell_margins(c, 140, 140, 200, 200)
    bp = c.paragraphs[0]
    bp.paragraph_format.space_after = Pt(2)
    b_title = bp.add_run("Didaktisk vejledning til læreren (Stilladseringen):")
    b_title.font.bold = True
    b_title.font.color.rgb = NAVY

    steps = [
        ("Trin 1: Makker-arbejde i appen (To om én skærm): ", "Kursisterne må IKKE sidde alene. Rolle A er 'Chauffør' (styrer appen og trykker på lyd). Rolle B er 'Sekretær/Ekko' (gentager udtale og noterer). Sammen læser de svarmulighederne og sætter kryds ved den rigtige overførte betydning."),
        ("Trin 2: Kobl livssituationer på talemåderne: ", "I grupper på 3-4 læser kursisterne autentiske livssituationer og diskuterer, hvilken talemåde der rammer plet, og hvorfor."),
        ("Trin 3: A/B Rollespil (Samtaletegnik & Udtale): ", "Kursisterne øver dialoger med faste samtalestrategier (aktiv lytning, empati, rådgivning) og fokus på hovedtryk i talemåderne."),
        ("Trin 4: Den Kulturelle Idiom-Børs: ", "Kursisterne bygger bro til deres egne modersmål og forklarer modersmålets talemåde på dansk.")
    ]
    for st_title, st_desc in steps:
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run("• " + st_title)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = DARK_GRAY
        r2 = p.add_run(st_desc)
        r2.font.size = Pt(9.5)

    doc.add_page_break()

    # =========================================================================
    # DEL 1: MAKKER-ARBEJDSARK: AFKODNING & BETYDNINGSVALG
    # =========================================================================
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("DEL 1: Makker-arbejdsark – Find den rigtige betydning")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = NAVY

    inst1 = doc.add_paragraph()
    inst1.paragraph_format.space_after = Pt(12)
    inst1.add_run(
        "Arbejd sammen to og to om én skærm i appen. Lyt til talemåden i appen, øv udtalen sammen, og læs de 4 svarmuligheder. "
        "Sæt kryds i boksen [X] ud for den korrekte forklaring."
    )
    inst1.runs[0].font.italic = True

    modules_data = [
        {
            "num": "Modul 1",
            "title": "At lykkes & Målrettet flid (Succes & Held)",
            "color": "E0F2FE",
            "items": [
                {
                    "idiom": "Blind høne finder også et korn",
                    "phonetic": "Udtale & tryk: [blind HØne finder også et KORN]",
                    "literal": "Bogstaveligt billede: En høne der ikke kan se, men tilfældigt finder mad på jorden.",
                    "options": [
                        ("At man kan være heldig at lykkes med noget, selvom man egentlig ikke er særlig dygtig eller forberedt.", True),
                        ("At en person, der arbejder i det stille, altid ender med at finde den største belønning.", False),
                        ("At man er så stædig, at man stjæler andres hårdt tjente resultater.", False),
                        ("At et dårligt syn gør, at man lægger mærke til de helt små detaljer på jorden.", False)
                    ]
                },
                {
                    "idiom": "At skyde papegøjen",
                    "phonetic": "Udtale & tryk: [at skyde papeGØJen]",
                    "literal": "Bogstaveligt billede: Gamle skydeselskaber hvor man skød efter en træfugl øverst på en mast.",
                    "options": [
                        ("At man er usædvanligt heldig og vinder den helt store gevinst (f.eks. drømmejobbet eller en god partner).", True),
                        ("At man er så utroligt præcis i sit arbejde, at man kan ramme et meget lille mål.", False),
                        ("At man ødelægger noget smukt og sjældent, bare fordi man vil vise sig frem for andre.", False),
                        ("At man tager på en lang jagt for at finde den eneste løsning på et svært problem.", False)
                    ]
                },
                {
                    "idiom": "At ramme plet",
                    "phonetic": "Udtale & tryk: [at ramme PLET]",
                    "literal": "Bogstaveligt billede: At ramme præcis midten af en målskive med en pil.",
                    "options": [
                        ("At man siger eller gør noget, der er fuldstændig rigtigt og præcist i situationen.", True),
                        ("At man vinder en stor sum penge i et meget svært og kompliceret spil.", False),
                        ("At man ved et uheld rammer et ømt punkt hos en anden og gør dem kede af det.", False),
                        ("At man skyder til måls og formår at træffe skivens kant.", False)
                    ]
                },
                {
                    "idiom": "At glide på skinner",
                    "phonetic": "Udtale & tryk: [at glide på SKINner]",
                    "literal": "Bogstaveligt billede: Et tog der kører let og stabilt fremad på jernbaneskinner.",
                    "options": [
                        ("At en proces eller opgave forløber helt planmæssigt, nemt og uden nogen form for problemer.", True),
                        ("At man har så meget energi, at man kan arbejde hurtigere end alle andre i teamet.", False),
                        ("At man sidder urokkeligt fast i en gammel rutine og nægter at prøve nye metoder.", False),
                        ("At man vælger at tage toget i stedet for bilen for at undgå trafikken på vejene.", False)
                    ]
                },
                {
                    "idiom": "At kaste frugt af sig",
                    "phonetic": "Udtale & tryk: [at kaste FRUGT af sig]",
                    "literal": "Bogstaveligt billede: Et frugttræ der efter lang tids pleje endelig bærer æbler.",
                    "options": [
                        ("At ens hårde arbejde endelig giver et godt, synligt og gavnligt resultat.", True),
                        ("At man er et meget generøst menneske, der elsker at forære mad og gaver til andre.", False),
                        ("At man spilder sine gode idéer på folk, som overhovedet ikke forstår dem.", False),
                        ("At man kaster med frugt, fordi man er utilfreds med høsten.", False)
                    ]
                }
            ]
        },
        {
            "num": "Modul 2",
            "title": "Mental & Handlingsretræte (Når man må trække sig)",
            "color": "FEF3C7",
            "items": [
                {
                    "idiom": "At stikke fingeren i jorden",
                    "phonetic": "Udtale & tryk: [at stikke FINgeren i JORden]",
                    "literal": "Bogstaveligt billede: Bonden mærker på jordens fugtighed og varme før såning.",
                    "options": [
                        ("At man er forsigtig og lige vurderer situationen og stemningen, før man handler.", True),
                        ("At man lytter ekstremt godt efter, hvad andre fortæller en, for at hjælpe dem.", False),
                        ("At man lukker sig fuldstændig inde i sig selv og nægter at tale med sine venner.", False),
                        ("At man arbejder i haven og gør sine hænder beskidte.", False)
                    ]
                },
                {
                    "idiom": "At stikke piben ind",
                    "phonetic": "Udtale & tryk: [at stikke PIben IND]",
                    "literal": "Bogstaveligt billede: Fiskeren eller soldaten gemmer piben væk for ikke at blive opdaget eller i modvind.",
                    "options": [
                        ("At man dæmper sig, bliver mindre påståelig og tier stille, fordi man indser, man tager fejl.", True),
                        ("At man kommer med en fantastisk idé, som løser problemet for hele gruppen.", False),
                        ("At man stjæler opmærksomheden fra en anden, netop som de skal til at tale.", False),
                        ("At man holder op med at ryge for at passe på sit helbred.", False)
                    ]
                },
                {
                    "idiom": "At trække i land",
                    "phonetic": "Udtale & tryk: [at trække i LAND]",
                    "literal": "Bogstaveligt billede: Sømanden ror båden tilbage mod havnen, når vejret bliver for hårdt.",
                    "options": [
                        ("At man ændrer mening eller trækker en tidligere påstand tilbage for at dæmpe en konflikt.", True),
                        ("At man redder et andet menneske op af vandet i en farlig situation.", False),
                        ("At man nægter at betale sin gæld tilbage til en ven, der har hjulpet en.", False),
                        ("At man fortøjrer sin båd solidt til molen efter en lang fisketur.", False)
                    ]
                },
                {
                    "idiom": "At slå bak",
                    "phonetic": "Udtale & tryk: [at slå BAK]",
                    "literal": "Bogstaveligt billede: Et skib eller lokomotiv der bakker med maskinen for at bremse op.",
                    "options": [
                        ("At man bremser op, fortryder sin handling og trækker sig tilbage fra det, man var i gang med.", True),
                        ("At man slår et stort og flot brød op, som man deler med hele familien.", False),
                        ("At man skubber til en anden person for at komme forrest i køen.", False),
                        ("At man bakker ud af en parkeringsplads med sin bil.", False)
                    ]
                },
                {
                    "idiom": "At kaste håndklædet i ringen",
                    "phonetic": "Udtale & tryk: [at kaste håndKLÆdet i RINGen]",
                    "literal": "Bogstaveligt billede: Boksetræneren smider håndklædet for at stoppe kampen og skåne bokseren.",
                    "options": [
                        ("At man giver helt op og erkender sit nederlag, fordi situationen er blevet for hård at klare.", True),
                        ("At man smider sine bekymringer væk og endelig begynder at nyde livet.", False),
                        ("At man kaster sit affald i naturen i stedet for at rydde ordentligt op.", False),
                        ("At man er færdig med at tage bad og smider håndklædet til vask.", False)
                    ]
                }
            ]
        },
        {
            "num": "Modul 3",
            "title": "Fejltrin & Kollaps (Når det går galt i livet)",
            "color": "FEE2E2",
            "items": [
                {
                    "idiom": "At træde i spinaten",
                    "phonetic": "Udtale & tryk: [at træde i spiNAten]",
                    "literal": "Bogstaveligt billede: At trampe rundt i den sarte køkkenhave og mase de fine grøntsager.",
                    "options": [
                        ("At man kommer til at sige eller gøre noget dumt, der skaber en pinlig og akavet situation.", True),
                        ("At man samler sunde urter og grøntsager i naturen til aftensmaden.", False),
                        ("At man ødelægger andres arbejde med vilje, fordi man er misundelig.", False),
                        ("At man får mudder på sine sko, når man går tur i skoven.", False)
                    ]
                },
                {
                    "idiom": "At slå et for stort brød op",
                    "phonetic": "Udtale & tryk: [at slå et for STORT BRØD op]",
                    "literal": "Bogstaveligt billede: At ælte så meget dej, at bageovnen ikke kan rumme brødet.",
                    "options": [
                        ("At man påtager sig en opgave eller et ansvar, der er alt for stort til, at man kan klare det.", True),
                        ("At man bager så lækkert et brød, at man kan invitere hele nabolaget på kaffe.", False),
                        ("At man spiser så meget mad, at man får ondt i maven og må hvile sig.", False),
                        ("At man lærer nye opskrifter på bagerkurset.", False)
                    ]
                },
                {
                    "idiom": "At gå i baglås",
                    "phonetic": "Udtale & tryk: [at gå i BAGlås]",
                    "literal": "Bogstaveligt billede: En lås der kiler sig fast, så nøglen slet ikke kan drejes.",
                    "options": [
                        ("At man fuldstændig mister evnen til at tænke klart eller handle på grund af stress eller pres.", True),
                        ("At man låser sin hoveddør så godt, at ingen tyve kan bryde ind.", False),
                        ("At man låser en person inde på et værelse under en konflikt.", False),
                        ("At man glemmer sine nøgler på arbejdspladsen og må vente på hjælp.", False)
                    ]
                },
                {
                    "idiom": "At stå med håret i postkassen",
                    "phonetic": "Udtale & tryk: [at stå med HÅret i POSTkassen]",
                    "literal": "Bogstaveligt billede: Man stikker hovedet for langt ind, klappen smækker i, og man hænger fast.",
                    "options": [
                        ("At man er havnet i en rigtig ærgerlig og pinlig knibe, hvor man står magtesløs tilbage som taberen.", True),
                        ("At man modtager et ubehageligt brev med en stor regning, man ikke kan betale.", False),
                        ("At man kigger efter post i blæsevejr, så håret flyver rundt.", False),
                        ("At man stjæler andres breve og bliver opdaget af naboen.", False)
                    ]
                },
                {
                    "idiom": "At gå ned med flaget",
                    "phonetic": "Udtale & tryk: [at gå NED med FLAget]",
                    "literal": "Bogstaveligt billede: Krigsskibet der synker med flaget hejst til det sidste.",
                    "options": [
                        ("At man bryder helt sammen af udmattelse, stress eller sygdom og ikke orker mere.", True),
                        ("At man sænker Dannebrog ved solnedgang som en pæn tradition.", False),
                        ("At man taber en fodboldkamp, selvom man kæmpede bravt.", False),
                        ("At man pakker havemøblerne væk for vinteren.", False)
                    ]
                }
            ]
        },
        {
            "num": "Modul 4",
            "title": "Starte på en frisk (Resiliens & Genrejsning)",
            "color": "DCFCE7",
            "items": [
                {
                    "idiom": "At rejse sig ved det træ, man er faldet ved",
                    "phonetic": "Udtale & tryk: [at rejse sig ved det TRÆ, man er FALdet ved]",
                    "literal": "Bogstaveligt billede: At snuble over en trærod i skoven, men tage fat i stammen for at komme op.",
                    "options": [
                        ("At man overvinder sit nederlag ved at prøve igen samme sted eller opgave, hvor det gik galt.", True),
                        ("At man planter et nyt træ i skoven for hvert problem, man overvinder.", False),
                        ("At man nægter at forlade det sted, hvor man kom til skade.", False),
                        ("At man fælder brænde efter en storm for at få varme til huset.", False)
                    ]
                },
                {
                    "idiom": "At komme op på hesten igen",
                    "phonetic": "Udtale & tryk: [at komme op på HESten igen]",
                    "literal": "Bogstaveligt billede: Rytteren der kastes af hesten, men straks stiger op igen for ikke at blive bange.",
                    "options": [
                        ("At man hurtigt prøver igen efter et nederlag, så man ikke når at miste modet.", True),
                        ("At man vinder en flot pokal i en ridekonkurrence efter lang træning.", False),
                        ("At man tvinger en hest til at løbe hurtigere, selvom den er træt.", False),
                        ("At man tager ud på en hyggelig ridetur i det gode forårsvejr.", False)
                    ]
                },
                {
                    "idiom": "At tage skeen i den anden hånd",
                    "phonetic": "Udtale & tryk: [at tage SKEen i den ANden hånd]",
                    "literal": "Bogstaveligt billede: At skifte greb om spiseredskabet for at få bedre fat.",
                    "options": [
                        ("At man pludselig ændrer sin måde at gøre tingene på og tager fat med ny beslutsomhed.", True),
                        ("At man hjælper et sygt familiemedlem med at spise deres suppe.", False),
                        ("At man smider bestikket på bordet i protest over maden.", False),
                        ("At man øver sig i at skrive med venstre hånd.", False)
                    ]
                },
                {
                    "idiom": "At smøge ærmerne op",
                    "phonetic": "Udtale & tryk: [at smøge ÆRmerne OP]",
                    "literal": "Bogstaveligt billede: Håndværkeren ruller skjorteærmerne op for ikke at svine dem til under hårdt arbejde.",
                    "options": [
                        ("At man gør sig klar til at tage hårdt fat og yde en stor, praktisk indsats uden at klage.", True),
                        ("At man tager sin varme trøje af, fordi man har fået det for varmt.", False),
                        ("At man syr sine gamle skjorter om for at spare penge på tøj.", False),
                        ("At man viser sine muskler for at imponere kollegerne.", False)
                    ]
                },
                {
                    "idiom": "Hvor der er hjerterum, er der husrum",
                    "phonetic": "Udtale & tryk: [hvor der er HJERterum, er der HUSrum]",
                    "literal": "Bogstaveligt billede: Et lille, beskedent hjem hvor gæster altid bydes indenfor i varmen.",
                    "options": [
                        ("Når man har et gæstfrit hjerte og empati, er der altid plads til flere, selvom pladsen er trang.", True),
                        ("At man skal bygge et kæmpe hus, før man kan invitere sin familie på besøg.", False),
                        ("At man kun åbner sit hjem for folk, man kender rigtig godt i forvejen.", False),
                        ("At man skal passe godt på sit helbred og hjerte i et gammelt hus.", False)
                    ]
                }
            ]
        }
    ]

    for mod in modules_data:
        # Module Header
        mp = doc.add_paragraph()
        mp.paragraph_format.space_before = Pt(12)
        mp.paragraph_format.space_after = Pt(4)
        m_run = mp.add_run(f"{mod['num']}: {mod['title']}")
        m_run.font.size = Pt(13)
        m_run.font.bold = True
        m_run.font.color.rgb = NAVY

        for idx, item in enumerate(mod['items'], 1):
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "FAFAFA")
            set_cell_margins(cell, 100, 100, 150, 150)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            it_name = cp.add_run(f"Talemåde {idx}: »{item['idiom']}«\n")
            it_name.font.bold = True
            it_name.font.size = Pt(11.5)
            it_name.font.color.rgb = NAVY

            it_ph = cp.add_run(f"🗣️  {item['phonetic']}\n")
            it_ph.font.size = Pt(9.5)
            it_ph.font.color.rgb = TEAL

            it_lit = cp.add_run(f"🖼️  {item['literal']}\n\n")
            it_lit.font.size = Pt(9.5)
            it_lit.font.italic = True
            it_lit.font.color.rgb = DARK_GRAY

            it_prompt = cp.add_run("Hvad betyder talemåden i overført betydning? (Sæt kryds):\n")
            it_prompt.font.bold = True
            it_prompt.font.size = Pt(9.5)

            for opt_text, is_correct in item['options']:
                op_p = cell.add_paragraph()
                op_p.paragraph_format.space_after = Pt(2)
                box_run = op_p.add_run("  [   ]  ")
                box_run.font.bold = True
                box_run.font.size = Pt(10)
                box_run.font.color.rgb = DARK_GRAY
                text_run = op_p.add_run(opt_text)
                text_run.font.size = Pt(9.5)

            # Note line
            np = cell.add_paragraph()
            np.paragraph_format.space_before = Pt(4)
            np.paragraph_format.space_after = Pt(2)
            n_run = np.add_run("Vores forklaring med egne ord: __________________________________________________")
            n_run.font.size = Pt(9)
            n_run.font.color.rgb = DARK_GRAY

            # Spacing between cards
            sp_p = doc.add_paragraph()
            sp_p.paragraph_format.space_after = Pt(4)

        doc.add_page_break()

    # =========================================================================
    # DEL 2: LIVSSITUATIONER (MATCH MED TALEMÅDER)
    # =========================================================================
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r2 = h2.add_run("DEL 2: Livssituationer – Kobl talemåderne på virkeligheden")
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = NAVY

    inst2 = doc.add_paragraph()
    inst2.paragraph_format.space_after = Pt(10)
    inst2.add_run(
        "Arbejd i grupper på 3-4 kursister. Læs situationerne højt for hinanden på skift. "
        "Diskutér i gruppen: Hvilken talemåde passer bedst til personens situation? "
        "Brug hjælpesætningerne til jeres samtale."
    )
    inst2.runs[0].font.italic = True

    # Scaffold box for group talk
    h_box = doc.add_table(rows=1, cols=1)
    h_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    hc = h_box.cell(0, 0)
    set_cell_background(hc, "F8FAFC")
    set_cell_margins(hc, 100, 100, 150, 150)
    hp = hc.paragraphs[0]
    hp.paragraph_format.space_after = Pt(2)
    h_title = hp.add_run("💬 Hjælpesætninger til gruppediskussionen:")
    h_title.font.bold = True
    h_title.font.size = Pt(10)
    h_title.font.color.rgb = TEAL

    phrases = [
        "»Jeg synes, at denne situation passer til talemåden [...], fordi personen...«",
        "»Hvad tænker I? Er I enige, eller kan man også bruge [anden talemåde]?«",
        "»Ja, det giver god mening, for han/hun indså jo, at opgaven var alt for svær.«"
    ]
    for ph in phrases:
        p = hc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("• " + ph)
        r.font.size = Pt(9.5)
        r.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    situations = [
        {
            "num": "Situation A",
            "title": "Rasmus og de mange ekstraopgaver på kontoret",
            "desc": "Rasmus har lige fået nyt job som koordinator. Han vil rigtig gerne imponere chefen og vise sit værd, så i løbet af sin første måned sagde han ja til tre store IT-projekter, meldte sig som formand for festudvalget og tilbød at oplære to nye elever. Nu sover han kun 3-4 timer om natten, har konstant hjertebanken og glemmer vigtige deadlines.",
            "match": "At slå et for stort brød op  /  At gå ned med flaget"
        },
        {
            "num": "Situation B",
            "title": "Middagen hos svigerfamilien",
            "desc": "Fatima var inviteret hjem til sin danske kærestes forældre for første gang. Stemningen var hyggelig, og maden var god. Men for at lave lidt sjov kom Fatima til at sige, at det maleri, der hang over sofaen, lignede noget, en børnehaveklasse havde tegnet. Der blev lammende tavshed ved bordet – for svigermoderen havde selv malet billedet! Fatima blev fuldstændig rød i hovedet.",
            "match": "At træde i spinaten  /  At stikke piben ind"
        },
        {
            "num": "Situation C",
            "title": "Den mislykkede køreprøve",
            "desc": "Mikael dumpede til sin praktiske køreprøve i morges, fordi han overså et ubetinget vigepligt-skilt i et vejkryds. Han var sønderknust, da han kom hjem, og havde bare lyst til at gemme sig under dynen. Men efter en time tog han en dyb indånding, ringede til sin kørelærer og sagde: 'Find den allerførste ledige prøvetid til mig, jeg prøver igen med det samme!'",
            "match": "Op på hesten igen  /  At rejse sig ved det træ, man er faldet ved"
        },
        {
            "num": "Situation D",
            "title": "Konflikten på lærerværelset",
            "desc": "Lars var meget uenig med ledelsens nye mødeplan og hævdede højlydt på personalemødet, at ingen lærere ville acceptere de nye regler. Men da inspektøren spurgte resten af lærerne, bakkede ingen Lars op. Lars kunne mærke, at han stod helt alene med sin vrede, så han tog ordet igen og sagde: 'Okay, jeg overdrev måske lidt. Vi kan godt prøve planen i to måneder først.'",
            "match": "At stikke piben ind  /  At trække i land"
        },
        {
            "num": "Situation E",
            "title": "Den uventede eksamenssucces",
            "desc": "Jonas havde overhovedet ikke fået læst op til sin biologi-eksamen, fordi han havde haft influenza hele ugen op til. Han trak et emne om fotosyntese – præcis den ene dokumentarfilm, han tilfældigt havde set på YouTube aftenen forinden! Sensor og censor gav ham et 10-tal for hans levende forklaring.",
            "match": "Blind høne finder også et korn"
        },
        {
            "num": "Situation F",
            "title": "Kusinens akutte bolignød",
            "desc": "Ayas kusine og hendes to små børn stod pludselig uden tag over hovedet på grund af en vandskade i deres lejlighed. Aya bor selv i en lille toværelses lejlighed på Nørrebro med sin mand. Alligevel sagde Aya straks: 'I skal selvfølgelig bo her hos os! Vi lægger madrasser ud på stuegulvet, det finder vi sagtens ud af.'",
            "match": "Hvor der er hjerterum, er der husrum"
        },
        {
            "num": "Situation G",
            "title": "De 50 afslag og drømmejobbet",
            "desc": "Efter at have sendt over 50 jobansøgninger uden held var Elena ved at miste alt håb. Men i går blev hun ringet op af Novo Nordisk, som tilbød hende en fast stilling som laboratorieleder med høj løn og fantastiske kolleger. Elena dansede rundt i hele stuen og kunne slet ikke stoppe med at smile.",
            "match": "At skyde papegøjen  /  At have svært ved at få armene ned"
        },
        {
            "num": "Situation H",
            "title": "Det stoppede byggeprojekt",
            "desc": "Kasper ville bygge en stor træterrasse helt alene i sin sommerferie. Efter at have købt materialer for 20.000 kr. fandt han ud af, at jorden under grunden var sumpet, og at terrassen ville synke i løbet af få uger. Kasper indså, at han måtte stoppe med det samme, før han spildte flere penge.",
            "match": "At slå bak  /  At stikke fingeren i jorden"
        }
    ]

    for s in situations:
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        sc = t.cell(0, 0)
        set_cell_background(sc, "FFFFFF")
        set_cell_margins(sc, 100, 100, 150, 150)

        sp = sc.paragraphs[0]
        sp.paragraph_format.space_after = Pt(2)
        s_title = sp.add_run(f"📌 {s['num']}: {s['title']}\n")
        s_title.font.bold = True
        s_title.font.size = Pt(11)
        s_title.font.color.rgb = NAVY

        s_desc = sp.add_run(s['desc'] + "\n\n")
        s_desc.font.size = Pt(9.5)
        s_desc.font.color.rgb = DARK_GRAY

        sp_m = sc.add_paragraph()
        sp_m.paragraph_format.space_after = Pt(2)
        m_lbl = sp_m.add_run("Hvilken talemåde passer her? ")
        m_lbl.font.bold = True
        m_lbl.font.size = Pt(9.5)
        sp_m.add_run("______________________________________________________________\n")

        sp_r = sc.add_paragraph()
        sp_r.paragraph_format.space_after = Pt(2)
        r_lbl = sp_r.add_run("Hvorfor? (Begrundelse): ")
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        sp_r.add_run("______________________________________________________________")

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    # =========================================================================
    # DEL 3: ROLLESPIL A OG B (SAMTALETEKIK & UDTALE)
    # =========================================================================
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    r3 = h3.add_run("DEL 3: Rollespil A & B – Udtale & Samtaleteknik")
    r3.font.size = Pt(16)
    r3.font.bold = True
    r3.font.color.rgb = NAVY

    inst3 = doc.add_paragraph()
    inst3.paragraph_format.space_after = Pt(10)
    inst3.add_run(
        "Arbejd sammen to og to. Først læser I manuskriptet højt sammen to gange (byt roller undervejs). "
        "Læg mærke til ordene med UNDERSTREGNING – her skal der være kraftigt tryk på dansk! "
        "Brug aktiv lytning og kropssprog."
    )
    inst3.runs[0].font.italic = True

    roleplays = [
        {
            "title": "Rollespil 1: Det pinlige møde (Spinaten & Piben ind)",
            "roles": "Person A = Den flove ven  •  Person B = Den forstående kollega",
            "dialogue": [
                ("A", "Hej Peter... Åh, jeg har det simpelthen så skidt i dag. Jeg kunne synke i jorden af skam!"),
                ("B", "Nå for søren, Mette! Hvad er der sket? Du ser helt bleg ud."),
                ("A", "Jeg var til middag hos min kærestes forældre i går. Og jeg kom til at sige højt, at deres nye sofa lignede noget fra et loppemarked... og det var svigermor, der selv havde valgt den!"),
                ("B", "Av, for den da! Der fik du da vist <u>trådt i spinaten</u> med begge ben!"),
                ("A", "Ja, fuldstændig! Svigermor kiggede så surt på mig, at jeg næsten ikke kunne trække vejret. Hvad pokker skal jeg gøre nu?"),
                ("B", "Jamen, du må simpelthen <u>stikke piben ind</u>, ringe til hende og sige pænt undskyld. Måske købe en flot buket blomster?"),
                ("A", "Du har ret. Jeg må sluge stoltheden og stikke piben ind med det samme. Tak for rådet, Peter!")
            ],
            "focus": "Udtalefokus: [træde i spiNAten] • [stikke PIben IND] • Aktiv lytning: 'Nå for søren!', 'Av, for den da!'"
        },
        {
            "title": "Rollespil 2: Den udbrændte kollega (For stort brød & Slå bak)",
            "roles": "Person A = Den overbebyrdede kollega  •  Person B = Den erfarne mentor",
            "dialogue": [
                ("A", "Hej Thomas... Har du fem minutter? Jeg tror, mit hoved er ved at eksplodere."),
                ("B", "Selvfølgelig har jeg tid. Sæt dig ned. Hvad trykker dig?"),
                ("A", "Jeg lovede chefen at færdiggøre tre store budgetter inden fredag, og samtidig meldte jeg mig til at arrangere kurset for de nye medarbejdere. Jeg har sovet fire timer hver nat i to uger, og jeg kan slet ikke tænke klart mere."),
                ("B", "Ved du hvad, min kære ven? Det lyder altså til, at du har fået <u>slået et alt for stort brød op</u>!"),
                ("A", "Ja, jeg ville så gerne vise, at jeg kan klare det hele... Men nu er jeg bange for at <u>gå helt ned med flaget</u>."),
                ("B", "Lige præcis. Og sundheden kommer først. Du bliver nødt til at <u>slå bak</u> lige nu. Gå ind til chefen, forklar situationen ærligt, og giv et af projekterne videre til mig."),
                ("A", "Tror du ikke, chefen bliver sur på mig?"),
                ("B", "Nej, overhovedet ikke. En god chef respekterer, at man kender sine grænser, før man knækker sammen.")
            ],
            "focus": "Udtalefokus: [slå et for STORT BRØD op] • [slå BAK] • Samtalestrategi: Empati og beroligelse."
        },
        {
            "title": "Rollespil 3: Nederlaget til eksamen (Op på hesten igen)",
            "roles": "Person A = Den skuffede studerende  •  Person B = Den opmuntrende studiekammerat",
            "dialogue": [
                ("A", "Jeg dumpede til modultesten i dag... Jeg er så uendeligt skuffet over mig selv."),
                ("B", "Ej, det er jeg virkelig ked af at høre! Men hvad skete der? Du havde jo øvet dig så meget?"),
                ("A", "Jeg <u>gik fuldstændig i baglås</u> under den mundtlige del. Klappen gik bare ned, og jeg kunne ikke huske et eneste ord på dansk."),
                ("B", "Hør her: Alle mennesker kan have en dårlig dag. Det betyder ikke, at du ikke kan tale dansk!"),
                ("A", "Lige nu føler jeg bare, at jeg aldrig lærer det. Jeg har lyst til at <u>kaste håndklædet i ringen</u>."),
                ("B", "Nej, hør nu her! Du må ikke give op. Du skal <u>op på hesten igen</u> med det samme! Vi to sætter os sammen i morgen, øver spørgsmålene, og så går du op til reeksamen med oprejst pande."),
                ("A", "Tak, fordi du tror på mig... Du har ret, jeg må <u>rejse mig ved det træ, jeg er faldet ved</u>.")
            ],
            "focus": "Udtalefokus: [gå i BAGlås] • [op på HESten igen] • Samtalestrategi: Opmuntring og modstandskraft."
        },
        {
            "title": "Rollespil 4: Drømmejobbet i hus (Skyde papegøjen & Svært ved at få armene ned)",
            "roles": "Person A = Den lykkelige jobsøger  •  Person B = Den begejstrede ægtefælle / ven",
            "dialogue": [
                ("A", "Gæt engang, hvad der lige er sket! Gæt!"),
                ("B", "Hold da op, dit smil går jo fra det ene øre til det andet! Ringede de fra hospitalet?"),
                ("A", "JA! Jeg fik jobbet som sygeplejerske! Fastansættelse, god løn og kun dagvagter!"),
                ("B", "Nej, hvor er det fantastisk! Du har jo simpelthen <u>skudt papegøjen</u>!"),
                ("A", "Ja, jeg har helt <u>svært ved at få armene ned</u>! Jeg hoppede rundt i hele haven, da chefen sagde tillykke."),
                ("B", "Det har du også fortjent efter al den flid! Dit hårde arbejde har endelig <u>kastet frugt af sig</u>!"),
                ("A", "I aften skal vi fejre det med god mad og kage!")
            ],
            "focus": "Udtalefokus: [skyde papeGØJen] • [få ARmene NED] • Samtalestrategi: Fælles glæde og fejring."
        }
    ]

    for rp in roleplays:
        rpt = doc.add_table(rows=1, cols=1)
        rpt.alignment = WD_TABLE_ALIGNMENT.CENTER
        rpc = rpt.cell(0, 0)
        set_cell_background(rpc, "F8FAFC")
        set_cell_margins(rpc, 100, 100, 150, 150)

        rp_p = rpc.paragraphs[0]
        rp_p.paragraph_format.space_after = Pt(2)
        rpt_title = rp_p.add_run(f"🎭 {rp['title']}\n")
        rpt_title.font.bold = True
        rpt_title.font.size = Pt(11)
        rpt_title.font.color.rgb = NAVY

        rpt_roles = rp_p.add_run(f"Roller: {rp['roles']}\n\n")
        rpt_roles.font.size = Pt(9.5)
        rpt_roles.font.italic = True
        rpt_roles.font.color.rgb = TEAL

        for spkr, line in rp['dialogue']:
            dp = rpc.add_paragraph()
            dp.paragraph_format.space_after = Pt(2)
            s_run = dp.add_run(f"{spkr}: ")
            s_run.font.bold = True
            s_run.font.size = Pt(9.5)
            s_run.font.color.rgb = NAVY if spkr == 'A' else TEAL

            # Handle underline tags
            parts = line.split("<u>")
            for idx, part in enumerate(parts):
                if idx == 0:
                    dp.add_run(part).font.size = Pt(9.5)
                else:
                    subparts = part.split("</u>")
                    u_run = dp.add_run(subparts[0])
                    u_run.font.underline = True
                    u_run.font.bold = True
                    u_run.font.size = Pt(9.5)
                    u_run.font.color.rgb = RGBColor(185, 28, 28) # Red emphasis
                    if len(subparts) > 1:
                        dp.add_run(subparts[1]).font.size = Pt(9.5)

        fp = rpc.add_paragraph()
        fp.paragraph_format.space_before = Pt(4)
        fp.paragraph_format.space_after = Pt(2)
        f_run = fp.add_run(rp['focus'])
        f_run.font.size = Pt(9)
        f_run.font.bold = True
        f_run.font.color.rgb = DARK_GRAY

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    # =========================================================================
    # DEL 4: DEN KULTURELLE IDIOM-BØRS
    # =========================================================================
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(4)
    r4 = h4.add_run("DEL 4: Den Kulturelle Idiom-Børs (Kontrastiv syntese)")
    r4.font.size = Pt(16)
    r4.font.bold = True
    r4.font.color.rgb = NAVY

    inst4 = doc.add_paragraph()
    inst4.paragraph_format.space_after = Pt(10)
    inst4.add_run(
        "Nu bygger vi bro mellem dansk kultur og jeres egne sprog. "
        "Menneskets grundvilkår er de samme over hele verden, men billederne i sproget er forskellige. "
        "Vælg én eller to danske talemåder og find ud af, hvordan man udtrykker det samme på dit modersmål. "
        "Forklar dit modersmåls talemåde på dansk!"
    )
    inst4.runs[0].font.italic = True

    cult_table = doc.add_table(rows=1, cols=1)
    cult_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cc = cult_table.cell(0, 0)
    set_cell_background(cc, "FFFBEB") # Warm light yellow
    set_cell_margins(cc, 120, 120, 150, 150)

    cp = cc.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    c_title = cp.add_run("🌍 Eksempel på en sammenligning:")
    c_title.font.bold = True
    c_title.font.size = Pt(10.5)
    c_title.font.color.rgb = RGBColor(180, 83, 9)

    ex_p = cc.add_paragraph()
    ex_p.paragraph_format.space_after = Pt(2)
    ex_run = ex_p.add_run(
        "Dansk talemåde: »At komme op på hesten igen« (eller »At rejse sig ved det træ, man er faldet ved«)\n"
        "På arabisk: »سقطة تعلّمك المشي« (Bogstaveligt: 'Faldet lærer dig at gå').\n"
        "Dansk forklaring: 'I mit land siger vi, at når et lille barn falder og slår sig, er det netop faldet, der lærer barnet at gå ordentligt. Man skal ikke være bange for at fejle, for det gør én stærkere!'"
    )
    ex_run.font.size = Pt(9.5)
    ex_run.font.italic = True
    ex_run.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Empty fill-in worksheet for students
    for i in range(1, 3):
        wt = doc.add_table(rows=1, cols=1)
        wt.alignment = WD_TABLE_ALIGNMENT.CENTER
        wc = wt.cell(0, 0)
        set_cell_background(wc, "FFFFFF")
        set_cell_margins(wc, 100, 100, 150, 150)

        wp = wc.paragraphs[0]
        wp.paragraph_format.space_after = Pt(2)
        w_title = wp.add_run(f"Min talemåde {i} fra mit hjemland:\n")
        w_title.font.bold = True
        w_title.font.size = Pt(11)
        w_title.font.color.rgb = NAVY

        fields = [
            ("Den danske talemåde jeg sammenligner med:", "________________________________________________"),
            ("Hvad siger man på mit modersmål?:", "________________________________________________"),
            ("Hvad betyder det helt bogstaveligt (ord for ord)?:", "________________________________________________"),
            ("Hvilket kulturelt billede bruges (dyr, natur, mad)?:", "________________________________________________"),
            ("Min forklaring på dansk til klassen:", "\n_________________________________________________________________________________\n_________________________________________________________________________________")
        ]
        for f_lbl, f_line in fields:
            fp = wc.add_paragraph()
            fp.paragraph_format.space_after = Pt(3)
            r1 = fp.add_run("• " + f_lbl + " ")
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = DARK_GRAY
            r2 = fp.add_run(f_line)
            r2.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Save document
    os.makedirs("public/downloads", exist_ok=True)
    out_path = "public/downloads/Workshop_Livets_Gang_Danske_Talemaader.docx"
    doc.save(out_path)
    print(f"Document successfully created at {out_path}")

if __name__ == "__main__":
    create_document()
